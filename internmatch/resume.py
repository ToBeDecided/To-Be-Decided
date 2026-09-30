"""Resume text extraction and rule-based scoring.

The score is a transparent rubric tuned for pre-law, business and humanities
internships. Each sub-score comes with the evidence behind it, and each
weakness comes with a concrete fix.
"""

from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from datetime import date

from .models import BulletFeedback, Profile, ResumeReport, SubScore
from .skills import CATEGORY_SIGNATURES, LANGUAGE_NAMES, extract_skills, skill_group

# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------


class ResumeParseError(ValueError):
    pass


# Formats macOS's built-in `textutil` can turn into plain text.
_TEXTUTIL_EXTS = (".doc", ".rtf", ".rtfd", ".odt", ".wordml", ".webarchive")


def extract_text(filename: str, data: bytes) -> tuple[str, int | None]:
    """Return (text, page_count) for a PDF, Word, Pages, RTF or plain-text resume."""
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        return _pdf_text(data)
    if name.endswith(".pages"):
        return _pages_text(data)
    if name.endswith(".docx") or (data[:2] == b"PK" and not name.endswith(_TEXTUTIL_EXTS)):
        return _docx_text(data), None
    if name.endswith(_TEXTUTIL_EXTS) or data[:5] == b"{\\rtf":
        text = _textutil_text(name, data)
        if text is not None:
            return text, None
        if data[:5] == b"{\\rtf":
            return rtf_to_text(data.decode("latin-1")), None
        raise ResumeParseError(
            f"{os.path.splitext(name)[1] or 'This'} files can only be read on a Mac. "
            "Save your resume as a PDF or .docx and try again."
        )
    try:
        return data.decode("utf-8"), None
    except UnicodeDecodeError:
        return data.decode("latin-1"), None


def _pdf_text(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise ResumeParseError(f"Couldn't read that PDF ({exc}).") from exc
    text = "\n".join(pages)
    if len(text.strip()) < 50:
        raise ResumeParseError(
            "That PDF has almost no extractable text (it may be a scanned image). "
            "Export your resume as a text-based PDF or upload the .docx."
        )
    return text, len(reader.pages)


def _docx_text(data: bytes) -> str:
    import docx

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of zip/xml errors
        raise ResumeParseError(f"Couldn't read that .docx ({exc}).") from exc
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        ppr = para._p.pPr
        is_list = "list" in (para.style.name or "").lower() or (ppr is not None and ppr.numPr is not None)
        lines.append(f"• {text}" if is_list and not _BULLET_RX.match(text) else text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text.strip())
    return "\n".join(lines)


def _pages_text(data: bytes) -> tuple[str, int | None]:
    """Apple Pages files are zip bundles; older ones carry a PDF preview we can read."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for member in ("QuickLook/Preview.pdf", "preview.pdf"):
                if member in zf.namelist():
                    return _pdf_text(zf.read(member))
    except zipfile.BadZipFile:
        pass
    raise ResumeParseError(
        "Pages files can't be read directly. In Pages, choose File > Export To > PDF, then upload the PDF."
    )


def _textutil_text(name: str, data: bytes) -> str | None:
    """Convert .doc/.rtf/... with macOS's built-in textutil. Returns None when it isn't available."""
    if sys.platform != "darwin" or not shutil.which("textutil"):
        return None
    suffix = os.path.splitext(name)[1] or ".rtf"
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, f"resume{suffix}")
        with open(path, "wb") as fh:
            fh.write(data)
        try:
            out = subprocess.run(["textutil", "-convert", "txt", "-stdout", "--", path],
                                 capture_output=True, timeout=30, check=True)
        except (subprocess.SubprocessError, OSError) as exc:
            raise ResumeParseError(f"Couldn't convert that {suffix} file ({exc}).") from exc
    return out.stdout.decode("utf-8", errors="replace")


def rtf_to_text(rtf: str) -> str:
    """Minimal RTF-to-text for platforms without textutil. Good enough for resumes."""
    out: list[str] = []
    skip_depth: list[int] = []
    depth = 0
    i = 0
    while i < len(rtf):
        ch = rtf[i]
        if ch == "{":
            depth += 1
            if rtf.startswith("{\\*", i) or re.match(r"\{\\(fonttbl|colortbl|stylesheet|info|pict)", rtf[i:i + 12]):
                skip_depth.append(depth)
            i += 1
        elif ch == "}":
            if skip_depth and skip_depth[-1] == depth:
                skip_depth.pop()
            depth -= 1
            i += 1
        elif ch == "\\":
            m = re.match(r"\\([a-z]+)(-?\d+)? ?|\\'([0-9a-f]{2})|\\(.)", rtf[i:], re.I)
            if not m:
                i += 1
                continue
            i += m.end()
            if skip_depth:
                continue
            word, num, hexa, sym = m.groups()
            if hexa:
                out.append(bytes([int(hexa, 16)]).decode("cp1252", errors="replace"))
            elif sym:
                out.append(sym if sym in "{}\\" else "")
            elif word in {"par", "line", "row"}:
                out.append("\n")
            elif word == "tab" or word == "cell":
                out.append("\t")
            elif word == "bullet":
                out.append("•")
            elif word == "u" and num:
                out.append(chr(int(num) % 65536))
                if i < len(rtf) and rtf[i] == "?":
                    i += 1
        else:
            if not skip_depth and ch not in "\r\n":
                out.append(ch)
            i += 1
    return re.sub(r"[ \t]+\n", "\n", "".join(out)).strip()


# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------

_BULLET_RX = re.compile(r"^\s*([•\u2022\u25cf\u25aa\u25e6\u2023\u2219\u00b7\u27a2\u2043\uf0b7\uf0a7\u25ba*–—-]|o\s)\s*")

SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "education": ("education", "academic background", "academics", "education & coursework", "study abroad",
                  "education and honors"),
    "experience": (
        "experience", "work experience", "professional experience", "employment", "internships",
        "work history", "relevant experience", "legal experience", "business experience", "industry experience",
        "internship experience", "professional & leadership experience", "employment history",
    ),
    "projects": ("projects", "selected projects", "consulting projects", "academic projects", "research projects"),
    "skills": ("skills", "skills & interests", "skills and interests", "skills, activities & interests",
               "skills & languages", "skills and languages", "languages", "language skills", "technical skills",
               "computer skills", "software", "core competencies", "additional information", "additional",
               "certifications & skills", "skills & certifications", "interests", "tools"),
    "leadership": ("leadership", "activities", "extracurriculars", "extracurricular activities", "involvement",
                   "organizations", "volunteer", "volunteering", "volunteer experience", "community service",
                   "leadership & activities", "leadership and activities", "leadership experience",
                   "leadership & involvement", "campus involvement", "community involvement", "clubs",
                   "memberships", "affiliations", "service", "activities & leadership"),
    "awards": ("awards", "honors", "achievements", "honors & awards", "awards & honors", "certifications",
               "certificates", "awards and honors", "scholarships", "honors and activities"),
    "coursework": ("coursework", "relevant coursework", "relevant courses", "selected coursework"),
    "summary": ("summary", "objective", "profile", "about me", "professional summary", "career objective"),
    "research": ("research", "research experience", "research & publications"),
    "publications": ("publications", "writing", "selected writing", "published work", "writing samples", "clips",
                     "selected publications", "articles"),
    "teaching": ("teaching", "teaching experience", "tutoring"),
}
_HEADER_LOOKUP = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}

ACTION_VERBS = frozenset(
    """
    accelerated accomplished achieved acquired adapted addressed administered advanced advised advocated analyzed
    applied argued arranged assembled assessed audited authored balanced boosted briefed budgeted built
    annotated archived calculated campaigned canvassed captured catalogued cataloged championed chaired coached coded
    coauthored cofounded documented examined gathered inventoried
    coled collaborated collected compiled completed composed conceived conducted consolidated constructed
    consulted contributed converted coordinated copyedited counseled created cultivated curated customized cut
    debated decreased defined delivered demonstrated designed determined developed devised digitized directed
    discovered drafted drove edited educated eliminated enabled engaged enhanced established evaluated executed
    expanded expedited facilitated filmed forecasted formulated fostered founded fundraised generated grew guided
    halved headed hosted identified implemented improved increased influenced initiated interpreted interviewed
    introduced investigated launched lectured led leveraged liaised lobbied maintained managed mapped marketed
    maximized measured mediated mentored merged minimized moderated modeled monitored motivated negotiated
    obtained operated optimized orchestrated organized oversaw partnered performed persuaded photographed
    pioneered piloted pitched planned prepared presented prioritized processed produced programmed promoted
    proofread proposed prospected published quantified raised ran recommended reconciled recruited redesigned
    reduced refined reorganized reported represented researched resolved restructured revamped reviewed revised
    rewrote saved scheduled scoped secured selected served shaped simplified sold solved spearheaded
    standardized started streamlined strengthened structured summarized supervised surveyed synthesized taught
    tracked trained transcribed transformed translated tutored unified upgraded valued verified volunteered
    welcomed won wrote
    """.split()
)
_PRESENT_TENSE = {"build", "design", "develop", "lead", "create", "implement", "write", "run", "own", "draft",
                  "research", "manage", "coordinate", "edit", "tutor", "teach", "advise", "analyze", "support",
                  "organize", "plan", "assist"}
_WEAK_OPENERS = (
    "responsible for", "helped", "help ", "worked on", "work on", "assisted", "assist ", "participated",
    "involved in", "tasked with", "duties included", "in charge of", "was ", "did ", "made ", "handled",
    "exposure to", "learned", "familiar with", "utilized", "attended", "member of",
)

_MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
_DATE_RANGE_RX = re.compile(
    rf"(?:{_MONTH}\s*'?\d{{2,4}}|\d{{1,2}}/\d{{2,4}}|(?:summer|fall|spring|winter)\s+\d{{4}}|\b(?:19|20)\d{{2}})"
    rf"\s*(?:-|–|—|to|until)\s*"
    rf"(?:{_MONTH}\s*'?\d{{2,4}}|\d{{1,2}}/\d{{2,4}}|\b(?:19|20)\d{{2}}|present|current|now|ongoing)",
    re.I,
)
_SEASON_RX = re.compile(r"\b(?:summer|fall|spring|winter)\s+(?:19|20)\d{2}\b", re.I)
_EMAIL_RX = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RX = re.compile(r"(?:\+?\d{1,2}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}\b")
_LINKEDIN_RX = re.compile(r"linkedin\.com/in/[\w%-]+", re.I)
_URL_RX = re.compile(r"(?:https?://)?(?:www\.)?[\w-]+\.(?:dev|io|me|com|net|org|app|site|xyz|co|art|studio|page)"
                     r"(?:/[\w./-]*)?\b", re.I)
_GPA_RX = re.compile(
    r"(?:gpa|g\.p\.a\.?|grade point average|cumulative)[^0-9\n]{0,20}([0-4]\.\d{1,3})(?:\s*/\s*(\d{1,2}(?:\.\d+)?))?",
    re.I,
)
_GPA_FALLBACK_RX = re.compile(r"\b([0-4]\.\d{1,2})\s*/\s*4\.0+\b")
_YEAR_RX = re.compile(r"\b(20\d{2})\b")
_PRONOUN_RX = re.compile(r"\b(I|me|my|My|we|our|Our)\b")

_LEADERSHIP_TITLE_RX = re.compile(
    r"\b(president|vice[- ]president|founder|co-?founder|captain|co-?captain|chair(?:person|woman|man)?|"
    r"co-?chair|director|editor[- ]in[- ]chief|managing editor|section editor|treasurer|secretary|"
    r"head of|team lead|officer|representative|student senator|delegate|coordinator|organizer|"
    r"resident advisor|\bra\b|orientation leader|ambassador|committee chair|board member|manager)\b",
    re.I,
)
_COMPETITION_RX = re.compile(
    r"mock trial|moot court|\bdebate\b|model u(nited )?n(ations)?|case competition|stock pitch|pitch competition|"
    r"ethics bowl|business plan competition|essay competition|writing contest|quiz bowl",
    re.I,
)
_PUBLICATION_RX = re.compile(
    r"\bpublished\b|student newspaper|student paper|staff writer|columnist|op-?eds?\b|journal article|"
    r"literary magazine|law review|undergraduate (law )?journal|byline|contributing writer|\bcolumn\b",
    re.I,
)
_HONORS_RX = re.compile(
    r"dean'?s list|honor society|phi beta kappa|phi alpha delta|pi sigma alpha|beta gamma sigma|"
    r"summa cum laude|magna cum laude|cum laude|honors (college|program|scholar)|scholarship|\bscholar\b|"
    r"presidential scholar|\bfellowship\b|award|prize|valedictorian|national merit",
    re.I,
)

_DEGREE_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("jd", re.compile(r"\bJ\.?\s?D\.?\b(?! Power)|\bJuris Doctor\b|\blaw school\b|\bSchool of Law\b", re.I)),
    ("phd", re.compile(r"\bPh\.?\s?D\b|\bDoctor of\b|\bDoctorate\b", re.I)),
    ("mba", re.compile(r"\bMBA\b")),
    ("master", re.compile(r"\bMasters?\b|\bMaster's\b|\bM\.S\.|\bM\.Sc\b|\bMSc\b|\bM\.A\.|\bMA (?:in|of)\b|"
                          r"\bMPP\b|\bMPA\b|\bMFA\b|\bM\.Ed\b|\bMEd\b")),
    ("bachelor", re.compile(r"\bBachelors?\b|\bBachelor's\b|\bB\.S\.|\bB\.A\.|\bBSc\b|\bB\.Sc\b|"
                            r"\bB\.?S\.? (?:in|of)\b|\bB\.?A\.? (?:in|of)\b|\bBBA\b|\bB\.B\.A\.|\bBFA\b|"
                            r"\bBS\b|\bBA\b|\bundergraduate\b", re.I)),
    ("associate", re.compile(r"\bAssociate'?s? (?:of|in|degree)\b|\bA\.S\.|\bA\.A\.", re.I)),
    ("high_school", re.compile(r"\bHigh School\b", re.I)),
)
_DEGREE_RANK = {"high_school": 0, "associate": 1, "bachelor": 2, "master": 3, "mba": 3, "jd": 3, "phd": 4}

MAJORS = (
    "Philosophy, Politics, and Economics", "Political Science", "International Relations", "International Affairs",
    "Public Policy", "Government", "Legal Studies", "Criminal Justice", "Criminology", "Economics",
    "Business Administration", "Business Management", "Finance", "Accounting", "Marketing", "Management",
    "Entrepreneurship", "Supply Chain Management", "Human Resources", "Real Estate", "Communications",
    "Communication Studies", "Journalism", "Media Studies", "Public Relations", "Advertising", "English",
    "Creative Writing", "Comparative Literature", "Literature", "History", "Art History", "Classics", "Philosophy",
    "Religious Studies", "Linguistics", "Anthropology", "Sociology", "Psychology", "Gender Studies",
    "African American Studies", "American Studies", "Latin American Studies", "East Asian Studies",
    "Middle Eastern Studies", "Global Studies", "Spanish", "French", "German", "Italian", "Chinese", "Japanese",
    "Music", "Theatre", "Theater", "Film Studies", "Film", "Fine Arts", "Studio Art", "Education", "Urban Studies",
    "Environmental Studies", "Cognitive Science", "Mathematics", "Statistics", "Data Science", "Computer Science",
)


@dataclass
class Bullet:
    text: str
    section: str
    words: int = 0
    action: bool = False
    weak_opener: str | None = None
    quantified: bool = False
    pronoun: bool = False

    def issues(self) -> list[str]:
        out = []
        if self.weak_opener:
            out.append(f"Starts with a weak phrase (\"{self.weak_opener.strip()}\"). Lead with what you did.")
        elif not self.action:
            out.append("Doesn't open with a strong action verb.")
        if not self.quantified:
            out.append("No number showing scale or impact (people served, pages, events, $, %).")
        if self.words > 35:
            out.append(f"Too long ({self.words} words). Aim for 1–2 lines.")
        elif self.words < 6:
            out.append("Too short to show impact.")
        if self.pronoun:
            out.append("Drop first-person pronouns (I/my/we).")
        return out


@dataclass
class ParsedResume:
    text: str
    lines: list[str]
    sections: dict[str, list[str]] = field(default_factory=dict)
    section_order: list[str] = field(default_factory=list)
    bullets: list[Bullet] = field(default_factory=list)
    skills: dict[str, int] = field(default_factory=dict)
    languages: list[str] = field(default_factory=list)
    contact: dict[str, bool] = field(default_factory=dict)
    gpa: float | None = None
    degree: str | None = None
    major: str | None = None
    grad_year: int | None = None
    roles: int = 0
    internships: int = 0
    projects: int = 0
    leadership_roles: int = 0
    competitions: int = 0
    research: bool = False
    publications: bool = False
    volunteering: bool = False
    teaching: bool = False
    study_abroad: bool = False
    honors: bool = False
    has_coursework: bool = False
    has_objective: bool = False
    word_count: int = 0
    pages: int | None = None


def _header_key(line: str) -> str | None:
    clean = re.sub(r"[^a-zA-Z&/, ]", "", line).strip().lower().rstrip(",")
    clean = re.sub(r"\s+", " ", clean)
    if not clean or len(clean) > 40:
        return None
    if clean in _HEADER_LOOKUP:
        return _HEADER_LOOKUP[clean]
    # Tolerate "Technical Skills & Interests"-style variants on short lines.
    words = clean.split()
    if len(words) <= 4 and (line.isupper() or line.rstrip().endswith(":")):
        for alias, key in sorted(_HEADER_LOOKUP.items(), key=lambda kv: -len(kv[0])):
            if alias in clean:
                return key
    return None


def _is_quantified(text: str) -> bool:
    stripped = _DATE_RANGE_RX.sub(" ", text)
    stripped = re.sub(r"\b(?:19|20)\d{2}\b", " ", stripped)
    return bool(re.search(r"\d", stripped)) or bool(
        re.search(r"\b(doubled|tripled|halved|dozens|hundreds|thousands|millions|one|two|three|four|five|six|"
                  r"seven|eight|nine|ten|twelve|fifteen|twenty)\b", stripped, re.I)
    )


def _analyze_bullet(text: str, section: str) -> Bullet:
    words = text.split()
    b = Bullet(text=text, section=section, words=len(words))
    first = re.sub(r"[^a-z]", "", words[0].lower()) if words else ""
    lower = text.lower()
    b.weak_opener = next((w for w in _WEAK_OPENERS if lower.startswith(w)), None)
    b.action = b.weak_opener is None and (
        first in ACTION_VERBS
        or first + "d" in ACTION_VERBS  # "Organize" -> "organized"
        or first + "ed" in ACTION_VERBS  # "Draft" -> "drafted"
        or (first.endswith("s") and (first[:-1] + "ed" in ACTION_VERBS or first[:-1] + "d" in ACTION_VERBS))
        or first in _PRESENT_TENSE
    )
    b.quantified = _is_quantified(text)
    b.pronoun = bool(_PRONOUN_RX.search(text))
    return b


def parse_resume(text: str, pages: int | None = None) -> ParsedResume:
    text = text.replace("\r", "\n").replace("\u00a0", " ")
    raw_lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.split("\n")]
    lines = [ln for ln in raw_lines if ln]
    pr = ParsedResume(text=text, lines=lines, pages=pages)
    pr.word_count = len(re.findall(r"\b\w+\b", text))

    # --- sections -------------------------------------------------------
    current = "header"
    sections: dict[str, list[str]] = {"header": []}
    order: list[str] = []
    for ln in lines:
        key = _header_key(ln)
        if key:
            current = key
            if key not in sections:
                order.append(key)
            sections.setdefault(current, [])
            if re.search(r"\bobjective\b", ln, re.I):
                pr.has_objective = True
            # Inline header like "Skills: Excel, Westlaw" keeps its content.
            if ":" in ln and ln.split(":", 1)[1].strip():
                sections[current].append(ln.split(":", 1)[1].strip())
            continue
        sections.setdefault(current, []).append(ln)
    pr.sections = sections
    pr.section_order = order

    # --- bullets (merge wrapped continuation lines) ---------------------
    bullet_sections = {"experience", "projects", "leadership", "research", "summary", "header", "awards",
                       "publications", "teaching"}
    body_sections = ("experience", "projects", "leadership", "research", "publications", "teaching")
    merged: list[tuple[str, str]] = []
    for sec, sec_lines in sections.items():
        for ln in sec_lines:
            m = _BULLET_RX.match(ln)
            if m and len(ln) > len(m.group(0)) + 2:
                merged.append((sec, ln[m.end():].strip()))
            elif merged and merged[-1][0] == sec and ln[:1].islower() and not _DATE_RANGE_RX.search(ln):
                merged[-1] = (sec, merged[-1][1] + " " + ln)
    if len(merged) < 3:
        # Few or no bullet glyphs survived extraction; treat sentence-like lines in body sections as bullets.
        fallback = [
            (sec, _BULLET_RX.sub("", ln, count=1))
            for sec in body_sections
            for ln in sections.get(sec, [])
            if len(ln.split()) >= 6 and not _DATE_RANGE_RX.search(ln)
        ]
        if len(fallback) > len(merged):
            merged = fallback
    pr.bullets = [_analyze_bullet(t, s) for s, t in merged if s in bullet_sections]

    # --- contact --------------------------------------------------------
    head = "\n".join(lines[:12])
    no_emails = _EMAIL_RX.sub(" ", text)  # "jo@example.com" is not a portfolio site
    urls = [u for u in _URL_RX.findall(no_emails) if not re.search(r"linkedin|github", u, re.I)]
    pr.contact = {
        "email": bool(_EMAIL_RX.search(text)),
        "phone": bool(_PHONE_RX.search(head) or _PHONE_RX.search(text)),
        "linkedin": bool(_LINKEDIN_RX.search(text) or re.search(r"\blinkedin\b", head, re.I)),
        "portfolio": bool(urls) or bool(re.search(r"\b(portfolio|writing samples?|clips)\b", head, re.I)),
    }

    # --- education ------------------------------------------------------
    edu_text = "\n".join(sections.get("education", [])) or text
    m = _GPA_RX.search(edu_text) or _GPA_RX.search(text)
    if m:
        value = float(m.group(1))
        scale = float(m.group(2)) if m.group(2) else 4.0
        if scale and scale != 4.0 and scale > value:
            value = value / scale * 4.0
        if 0 < value <= 4.3:
            pr.gpa = round(min(value, 4.0), 2)
    else:
        m2 = _GPA_FALLBACK_RX.search(edu_text)
        if m2:
            pr.gpa = float(m2.group(1))

    found = [lvl for lvl, rx in _DEGREE_PATTERNS if rx.search(edu_text)]
    if found:
        pr.degree = max(found, key=lambda d: _DEGREE_RANK[d])
    majors = [mj for mj in MAJORS if re.search(rf"\b{re.escape(mj)}\b", edu_text)]
    if majors:
        # Prefer the major that appears first in the education section.
        pr.major = min(majors, key=lambda mj: edu_text.find(mj))

    this_year = date.today().year
    years = [int(y) for y in _YEAR_RX.findall(edu_text) if this_year - 6 <= int(y) <= this_year + 7]
    grad_hint = re.search(r"(?:expected|graduat\w*|class of|anticipated)[^0-9\n]{0,20}(20\d{2})", edu_text, re.I)
    if grad_hint:
        pr.grad_year = int(grad_hint.group(1))
    elif years:
        pr.grad_year = max(years)
    pr.has_coursework = "coursework" in sections or bool(re.search(r"coursework|relevant courses", text, re.I))
    pr.study_abroad = bool(re.search(r"study abroad|studied abroad|exchange program|semester abroad", text, re.I))

    # --- experience -----------------------------------------------------
    exp_lines = sections.get("experience", [])
    role_lines = [ln for ln in exp_lines if _DATE_RANGE_RX.search(ln) or _SEASON_RX.search(ln)]
    pr.roles = len(role_lines)
    exp_blob = "\n".join(exp_lines)
    pr.internships = len(re.findall(r"\bintern(?:ship)?\b|\bco-?op\b|\bsummer analyst\b|\bextern(?:ship)?\b",
                                    exp_blob, re.I))
    pr.internships = min(pr.internships, pr.roles) if pr.roles else min(pr.internships, 2)
    if pr.roles == 0 and exp_lines:
        # Headers without parsable dates: count non-bullet title-ish lines.
        pr.roles = min(4, sum(1 for ln in exp_lines if not _BULLET_RX.match(ln) and len(ln.split()) <= 12
                              and ln[:1].isupper()) // 2)

    proj_lines = sections.get("projects", [])
    proj_headers = [ln for ln in proj_lines
                    if not _BULLET_RX.match(ln) and ln[:1].isupper() and len(ln.split()) <= 14]
    proj_bullets = sum(1 for b in pr.bullets if b.section == "projects")
    pr.projects = min(6, max(len(proj_headers), (proj_bullets + 2) // 3 if proj_bullets else 0))

    # Leadership titles are counted on non-bullet lines (the role headers), anywhere outside education.
    header_lines = [ln for sec, sec_lines in sections.items() if sec not in {"education", "skills"}
                    for ln in sec_lines if not _BULLET_RX.match(ln)]
    pr.leadership_roles = min(4, sum(1 for ln in header_lines if _LEADERSHIP_TITLE_RX.search(ln)))
    pr.competitions = min(3, len({m.group(0).lower() for m in _COMPETITION_RX.finditer(text)}))

    lower = text.lower()
    pr.research = "research" in sections or bool(
        re.search(r"research (assistant|intern|fellow)|undergraduate research|honors thesis|senior thesis|"
                  r"independent research|research project", lower))
    pr.publications = "publications" in sections or bool(_PUBLICATION_RX.search(text))
    pr.volunteering = bool(re.search(r"volunteer|community service|americorps|pro bono", lower))
    pr.teaching = "teaching" in sections or bool(
        re.search(r"\btutor|teaching assistant|writing center|\bmentor|camp counselor|\binstructor\b", lower))
    pr.honors = "awards" in sections or bool(_HONORS_RX.search(text))

    pr.skills = extract_skills(text)
    pr.languages = sorted(s for s in pr.skills if s in LANGUAGE_NAMES)
    return pr


# ---------------------------------------------------------------------------
# Scoring
# ---------------------------------------------------------------------------

LEVELS = ("Freshman", "Sophomore", "Junior", "Senior", "Graduate student", "Law student", "Recent graduate")


def school_level(grad_year: int | None, degree: str | None, today: date | None = None) -> str:
    today = today or date.today()
    if degree == "jd":
        return "Law student"
    if degree in {"master", "phd", "mba"}:
        return "Graduate student"
    if not grad_year:
        return "Junior"  # the most common internship applicant; neutral default
    school_year_end = today.year + 1 if today.month >= 8 else today.year
    years_left = grad_year - school_year_end
    if years_left < 0:
        return "Recent graduate"
    return {0: "Senior", 1: "Junior", 2: "Sophomore"}.get(years_left, "Freshman")


_EXPECTED_EXPERIENCE = {"Freshman": 30, "Sophomore": 50, "Junior": 75, "Senior": 90, "Graduate student": 90,
                        "Law student": 85, "Recent graduate": 95}


def category_affinity(skills: set[str]) -> dict[str, float]:
    """0-1 fit for each category: the weighted share of its signature skills you have.

    Having half of a category's weighted signature is treated as a full fit; nobody lists everything.
    """
    out = {}
    for cat, sig in CATEGORY_SIGNATURES.items():
        total = sum(sig.values())
        got = sum(w for s, w in sig.items() if s in skills)
        out[cat] = min(1.0, got / (0.5 * total))
    return out


def _gpa_points(gpa: float | None) -> int:
    if gpa is None:
        return 55
    for cutoff, pts in ((3.8, 100), (3.6, 90), (3.4, 80), (3.2, 68), (3.0, 58), (2.7, 45)):
        if gpa >= cutoff:
            return pts
    return 35


def _grade(score: int) -> str:
    if score >= 88:
        return "A"
    if score >= 80:
        return "A-"
    if score >= 74:
        return "B+"
    if score >= 68:
        return "B"
    if score >= 62:
        return "B-"
    if score >= 55:
        return "C+"
    if score >= 48:
        return "C"
    return "D"


_GPA_FOCUSED = {"Legal", "Finance & Accounting", "Consulting & Business", "Government & Policy"}
_PORTFOLIO_FOCUSED = {"Media & Writing", "Marketing & Communications", "Arts & Culture"}


def score_resume(pr: ParsedResume, profile: Profile | None = None) -> ResumeReport:
    from .skills import categories_for  # local import keeps the module import graph simple

    profile = profile or Profile()
    strengths: list[tuple[int, str]] = []
    fixes: list[tuple[int, str]] = []  # (priority, message): lower priority number = more important

    gpa = profile.gpa if profile.gpa is not None else pr.gpa
    degree = profile.degree_level or pr.degree or "bachelor"
    grad_year = profile.grad_year or pr.grad_year
    level = school_level(grad_year, degree)
    skills = set(pr.skills) | {s for s in profile.extra_skills if s}
    languages = [s for s in skills if s in LANGUAGE_NAMES]
    aff = category_affinity(skills)
    targets = categories_for(profile.target_tracks, profile.target_categories)
    explicit_targets = bool(profile.target_tracks or profile.target_categories)
    focus = targets if explicit_targets else {max(aff, key=aff.get)}

    # --- impact & writing ------------------------------------------------
    bullets = pr.bullets
    nb = len(bullets)
    if nb:
        action = sum(b.action for b in bullets)
        quant = sum(b.quantified for b in bullets)
        weak = sum(bool(b.weak_opener) for b in bullets)
        long_ = sum(b.words > 35 for b in bullets)
        short = sum(b.words < 6 for b in bullets)
        pron = sum(b.pronoun for b in bullets)
        length_ok = nb - long_ - short
        # Numbers matter less in humanities roles than in finance, but "12 articles" or "40 clients" still help.
        impact = 100 * (
            0.4 * action / nb + 0.3 * min(1.0, (quant / nb) / 0.45) + 0.2 * length_ok / nb + 0.1 * (1 - weak / nb)
        ) - min(10, 3 * pron)
        impact_detail = f"{action}/{nb} bullets start with action verbs, {quant}/{nb} include numbers"
        if action / nb >= 0.75:
            strengths.append((2, f"Strong writing: {action} of {nb} bullets open with an action verb."))
        elif action / nb < 0.6:
            fixes.append((2, f"{nb - action} of {nb} bullets don't open with a strong action verb. Start each with "
                             "a verb like Drafted, Researched, Organized, Analyzed or Led."))
        if quant / nb >= 0.45:
            strengths.append((1, f"Results are concrete: {quant} of {nb} bullets include a number."))
        else:
            have = f"Only {quant} of {nb} bullets include" if quant else f"None of your {nb} bullets include"
            fixes.append((1, f"{have} a number. Add scale where you can: people served, articles written, "
                             "cases or documents reviewed, dollars raised, events run, % growth."))
        if weak:
            fixes.append((3, f"{weak} bullet(s) start with passive phrases like \"Responsible for\" or \"Assisted "
                             "with\". Say what you did and what changed because of it."))
        if long_:
            fixes.append((5, f"{long_} bullet(s) run past 35 words. Keep each to one or two lines."))
        if pron:
            fixes.append((6, "Remove first-person pronouns (I, my, we) from bullets."))
    else:
        impact = 20
        impact_detail = "No bullet points detected"
        fixes.append((1, "We couldn't find bullet points. Describe each role and activity with 2–4 bullets that "
                         "start with an action verb and end with a concrete result."))

    # --- experience -------------------------------------------------------
    points = (
        pr.internships * 30
        + min(4, max(0, pr.roles - pr.internships)) * 15
        + pr.leadership_roles * 8
        + pr.competitions * 8
        + pr.projects * 6
        + (15 if pr.research else 0)
        + (10 if pr.publications else 0)
        + (6 if pr.teaching else 0)
        + (6 if pr.volunteering else 0)
        + (5 if pr.study_abroad else 0)
    )
    expected = _EXPECTED_EXPERIENCE[level]
    experience = max(10.0, min(100.0, 100 * points / expected))
    exp_bits = []
    if pr.roles:
        exp_bits.append(f"{pr.roles} role(s)" + (f", {pr.internships} internship(s)" if pr.internships else ""))
    if pr.leadership_roles:
        exp_bits.append(f"{pr.leadership_roles} leadership position(s)")
    if pr.competitions:
        exp_bits.append("competitions")
    for flag, label in ((pr.research, "research"), (pr.publications, "publications"), (pr.teaching, "teaching"),
                        (pr.volunteering, "service"), (pr.study_abroad, "study abroad")):
        if flag:
            exp_bits.append(label)
    exp_detail = (", ".join(exp_bits) or "little experience detected") + f" (compared with a typical {level.lower()})"
    if pr.internships:
        strengths.append((0, f"Prior internship experience ({pr.internships}). This is the strongest signal "
                             "for recruiters."))
    if pr.leadership_roles >= 2:
        strengths.append((3, f"{pr.leadership_roles} leadership positions. Law schools and employers in these "
                             "fields weigh leadership heavily."))
    if pr.publications and focus & {"Media & Writing", "Legal", "Government & Policy", "Arts & Culture"}:
        strengths.append((3, "Published writing. Great evidence for writing-heavy roles; keep a couple of clips "
                             "ready to send."))
    if experience < 60:
        if not pr.leadership_roles and "leadership" not in pr.sections:
            fixes.append((0, f"For a {level.lower()}, the resume is light on experience. Add a Leadership & "
                             "Activities section: club roles, student government, mock trial, the campus paper, "
                             "volunteering and part-time jobs all count."))
        else:
            fixes.append((1, "Add more experience: a campus job, research assistantship, volunteer role or "
                             "part-time work shows responsibility even when it isn't in your target field."))
    if "Legal" in focus and not ({"Mock Trial", "Moot Court", "Debate", "Legal Office Experience", "Legal Research"}
                                 & skills):
        fixes.append((3, "For pre-law roles, show legal exposure: mock trial or debate, a law-related course, "
                         "legal research, or volunteering at a legal aid clinic or court."))
    if "Finance & Accounting" in focus and not ({"Financial Modeling", "Valuation", "Equity Research",
                                                 "Accounting", "Financial Analysis"} & skills):
        fixes.append((3, "For finance roles, show technical interest: an investment club or stock pitch, "
                         "accounting coursework, or a financial modeling course (list the tools you used)."))

    # --- skills -----------------------------------------------------------
    best_cat, best_aff = max(aff.items(), key=lambda kv: kv[1])
    focus_aff = max(aff[c] for c in focus)
    breadth = min(1.0, len(skills) / 12)
    skills_score = 100 * (0.35 * breadth + 0.45 * focus_aff + 0.1 * ("skills" in pr.sections)
                          + 0.1 * bool(languages))
    skills_detail = f"{len(skills)} skills detected; strongest fit: {best_cat}"
    if languages:
        strengths.append((4, f"Foreign language ability ({', '.join(sorted(languages))}). A real differentiator in "
                             "law, policy, nonprofit and international roles."))
    if len(skills) >= 12:
        strengths.append((5, f"Well-rounded skill set ({len(skills)} skills detected)."))
    elif len(skills) < 7:
        fixes.append((2, f"Only {len(skills)} skills detected. List the software (Excel, PowerPoint, Westlaw, "
                         "Canva, Adobe…), research methods and languages you actually use in a Skills & Interests "
                         "section. Applicant tracking systems match on these keywords."))
    if "skills" not in pr.sections:
        fixes.append((3, "Add a \"Skills & Interests\" (or \"Skills & Languages\") section at the bottom."))

    # --- academics --------------------------------------------------------
    academics = 0.7 * _gpa_points(gpa) + 0.15 * (100 if pr.honors else 40) + 0.15 * (100 if pr.has_coursework else 50)
    acad_detail = f"GPA {gpa:.2f}" if gpa is not None else "GPA not listed"
    if pr.major:
        acad_detail += f", {pr.major}"
    if pr.honors:
        acad_detail += ", honors"
    gpa_matters = bool(focus & _GPA_FOCUSED)
    if gpa is None:
        fixes.append((2 if gpa_matters else 4,
                      "No GPA found. If it's 3.3 or higher, list it. Law firms, banks and consulting firms "
                      "screen on GPA, and pre-law advisors will ask." if gpa_matters else
                      "No GPA found. If it's 3.3 or higher, list it."))
    elif gpa >= 3.6:
        strengths.append((4, f"Strong GPA ({gpa:.2f})."))
    elif gpa < 3.0:
        fixes.append((6, "A GPA under 3.0 can hurt automated screens. Consider leaving it off and leaning on "
                         "experience and major GPA (if higher)."))
    elif gpa < 3.5 and gpa_matters:
        fixes.append((5, "For finance, consulting and law-firm roles a 3.5+ is the usual screen. If your major "
                         "GPA is higher, list it alongside the cumulative one."))
    if pr.honors:
        strengths.append((6, "Academic honors listed (Dean's List, scholarships or honor societies)."))
    if not pr.has_coursework and level in {"Freshman", "Sophomore"}:
        fixes.append((5, "Add a line of relevant coursework (e.g. Constitutional Law, Financial Accounting, "
                         "Research Methods) to show preparation for your target field."))

    # --- format & completeness ---------------------------------------------
    fmt = 0.0
    for sec in ("education", "skills"):
        fmt += 18 if sec in pr.sections else 0
    fmt += 18 if ({"experience", "leadership", "projects"} & set(pr.sections)) else 0
    fmt += 10 * pr.contact["email"] + 6 * pr.contact["phone"] + 10 * pr.contact["linkedin"]
    too_long = (pr.pages or 0) > 1 or (pr.pages is None and pr.word_count > 850)
    too_short = pr.word_count < 250
    a_bit_short = not too_short and pr.word_count < 380
    fmt += 5 if too_long or too_short else 12 if a_bit_short else 20
    edu_first = bool(pr.section_order) and (pr.section_order[0] == "education" or
                                            (pr.section_order[0] == "summary" and pr.section_order[1:2] == ["education"]))
    if "education" not in pr.sections:
        fixes.append((2, "Add an Education section with school, degree, major, expected graduation date and GPA."))
    elif not edu_first and level not in {"Recent graduate"}:
        fmt -= 6
        fixes.append((4, "Move Education to the top. For students, law, finance and consulting recruiters expect "
                         "it first."))
    if pr.has_objective:
        fmt -= 4
        fixes.append((6, "Drop the Objective statement. It takes space from your experience and rarely helps."))
    if not pr.contact["email"]:
        fixes.append((1, "No email address found. Put your contact info at the top."))
    if not pr.contact["linkedin"]:
        fixes.append((5, "Add your LinkedIn URL to the header. Business and policy recruiters look you up."))
    if focus & _PORTFOLIO_FOCUSED and not pr.contact["portfolio"]:
        fixes.append((5, "Add a link to a portfolio or writing samples (a simple site or Google Drive folder). "
                         "Media, marketing and arts employers ask for clips."))
    if too_long and level not in {"Graduate student", "Law student"}:
        fixes.append((3, "Keep it to one page. For internships, one tight page beats two."))
    if too_short:
        fixes.append((2, f"Your resume is sparse ({pr.word_count} words). Aim for a full page (~400–650 words)."))
    elif a_bit_short:
        fixes.append((5, f"There's room for more detail ({pr.word_count} words). A full page is ~400–650 words: "
                         "add an activity, coursework, or another bullet on your biggest result."))
    fmt = max(0.0, min(100.0, fmt))
    if fmt >= 90:
        strengths.append((6, "Well-structured: Education first, all the key sections and contact details present."))

    subscores = [
        SubScore(key="impact", label="Impact & writing", score=round(impact), weight=0.25, detail=impact_detail),
        SubScore(key="experience", label="Experience & involvement", score=round(experience), weight=0.25,
                 detail=exp_detail),
        SubScore(key="skills", label="Skills & languages", score=round(skills_score), weight=0.15,
                 detail=skills_detail),
        SubScore(key="academics", label="Academics", score=round(academics), weight=0.15, detail=acad_detail),
        SubScore(key="format", label="Format & completeness", score=round(fmt), weight=0.20,
                 detail=f"{len([s for s in pr.sections if s != 'header'])} sections, "
                        f"{pr.pages or '?'} page(s), {pr.word_count} words"),
    ]
    for s in subscores:
        s.score = max(0, min(100, s.score))
    overall = round(sum(s.score * s.weight for s in subscores))

    grouped: dict[str, list[str]] = {}
    for name in sorted(skills, key=lambda n: (-pr.skills.get(n, 0), n.lower())):
        grouped.setdefault(skill_group(name), []).append(name)

    weak_bullets = sorted((b for b in bullets if b.issues()), key=lambda b: (-len(b.issues()), -b.words))[:6]

    return ResumeReport(
        overall=overall,
        grade=_grade(overall),
        subscores=subscores,
        strengths=[m for _, m in sorted(strengths)][:6],
        improvements=[m for _, m in sorted(fixes)][:12],
        skills=grouped,
        skill_count=len(skills),
        sections=[s for s in pr.sections if s != "header"],
        contact=pr.contact,
        stats={
            "bullets": nb,
            "quantified_bullets": sum(b.quantified for b in bullets),
            "action_bullets": sum(b.action for b in bullets),
            "word_count": pr.word_count,
            "pages": pr.pages,
            "roles": pr.roles,
            "internships": pr.internships,
            "leadership_roles": pr.leadership_roles,
            "projects": pr.projects,
        },
        detected={
            "gpa": pr.gpa,
            "grad_year": pr.grad_year,
            "degree": pr.degree,
            "major": pr.major,
            "level": level,
            "languages": ", ".join(pr.languages) or None,
        },
        weak_bullets=[BulletFeedback(text=b.text, issues=b.issues()) for b in weak_bullets],
    )
