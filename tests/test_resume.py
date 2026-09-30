import io
import zipfile
from datetime import date

import pytest

from internmatch.models import Profile
from internmatch.resume import (
    ResumeParseError,
    extract_text,
    parse_resume,
    rtf_to_text,
    school_level,
    score_resume,
)

from .conftest import FIXTURES


def test_parse_prelaw_resume(prelaw_text):
    pr = parse_resume(prelaw_text)
    assert pr.section_order[0] == "education"
    assert {"education", "experience", "leadership", "skills"} <= set(pr.sections)
    assert pr.contact == {"email": True, "phone": True, "linkedin": True, "portfolio": False}
    assert (pr.gpa, pr.degree, pr.major, pr.grad_year) == (3.82, "bachelor", "Political Science", 2028)
    assert pr.roles == 3 and pr.internships == 2
    assert pr.leadership_roles >= 1 and pr.competitions >= 1
    assert pr.research and pr.publications and pr.volunteering and pr.honors
    assert pr.languages == ["Portuguese", "Spanish"]
    assert {"Westlaw", "Legal Research", "Mock Trial", "Constituent Services"} <= set(pr.skills)


@pytest.mark.parametrize("name", ["prelaw", "business", "humanities"])
def test_strong_resumes_score_well(name):
    report = score_resume(parse_resume((FIXTURES / f"{name}_resume.txt").read_text()))
    assert report.overall >= 85
    assert report.grade.startswith("A")
    assert all(0 <= s.score <= 100 for s in report.subscores)
    assert abs(sum(s.weight for s in report.subscores) - 1) < 1e-9


def test_weak_resume_gets_field_specific_fixes(weak_text):
    report = score_resume(parse_resume(weak_text))
    assert report.overall < 45
    text = " ".join(report.improvements)
    assert "Leadership & Activities" in text or "more experience" in text
    assert "number" in text and "action verb" in text
    assert "Move Education to the top" in text
    assert "Objective" in text
    assert "GitHub" not in text  # not a tech resume
    assert report.detected["degree"] == "associate"


def test_targets_change_advice(humanities_text, business_text):
    finance = score_resume(parse_resume(humanities_text), Profile(target_tracks=["Business"],
                                                                   target_categories=["Finance & Accounting"]))
    assert any("finance roles" in fix for fix in finance.improvements)
    media = score_resume(parse_resume(business_text), Profile(target_categories=["Media & Writing"]))
    assert any("portfolio or writing samples" in fix for fix in media.improvements)


def test_profile_overrides_detected_values(prelaw_text):
    report = score_resume(parse_resume(prelaw_text), Profile(gpa=3.1, grad_year=2030))
    academics = next(s for s in report.subscores if s.key == "academics")
    assert "3.10" in academics.detail
    assert report.detected["level"] == "Freshman"
    assert report.detected["gpa"] == 3.82


def test_school_level():
    today = date(2026, 9, 30)
    assert school_level(2027, "bachelor", today) == "Senior"
    assert school_level(2028, "bachelor", today) == "Junior"
    assert school_level(2029, "bachelor", today) == "Sophomore"
    assert school_level(2030, "bachelor", today) == "Freshman"
    assert school_level(2026, "bachelor", today) == "Recent graduate"
    assert school_level(2028, "jd", today) == "Law student"
    assert school_level(2027, "master", today) == "Graduate student"


def test_extract_pdf():
    text, pages = extract_text("resume.pdf", (FIXTURES / "prelaw_resume.pdf").read_bytes())
    assert pages == 1 and "Maya Thompson" in text
    report = score_resume(parse_resume(text, pages=pages))
    assert report.stats["pages"] == 1 and report.overall >= 85


def test_extract_docx_marks_list_paragraphs_as_bullets():
    import docx

    doc = docx.Document()
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Legal Intern, Acme LLP  Jun 2026 – Aug 2026")
    doc.add_paragraph("Drafted 12 memos on employment law using Westlaw", style="List Bullet")
    doc.add_paragraph("Organized 3,000 pages of discovery for trial", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    text, pages = extract_text("resume.docx", buf.getvalue())
    assert pages is None and "• Drafted 12 memos" in text
    pr = parse_resume(text)
    assert len(pr.bullets) == 2 and all(b.quantified and b.action for b in pr.bullets)


def test_rtf_resume():
    rtf = (r"{\rtf1\ansi{\fonttbl{\f0 Helvetica;}}\f0 Jordan Lee\par "
           r"\bullet  Drafted 12 policy memos\par Caf\'e9 manager\par \u8212? Spanish}")
    text = rtf_to_text(rtf)
    assert "Jordan Lee" in text and "Drafted 12 policy memos" in text and "Café" in text and "—" in text
    assert "Helvetica" not in text
    out, _ = extract_text("resume.rtf", rtf.encode("latin-1"))
    assert "Drafted 12 policy memos" in out


def test_pages_files():
    with open(FIXTURES / "prelaw_resume.pdf", "rb") as fh:
        pdf = fh.read()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("QuickLook/Preview.pdf", pdf)
    text, pages = extract_text("resume.pages", buf.getvalue())
    assert "Maya Thompson" in text and pages == 1

    modern = io.BytesIO()
    with zipfile.ZipFile(modern, "w") as zf:
        zf.writestr("Index/Document.iwa", b"\x00binary")
    with pytest.raises(ResumeParseError, match="Export To"):
        extract_text("resume.pages", modern.getvalue())


def test_doc_files_need_textutil(monkeypatch):
    monkeypatch.setattr("sys.platform", "linux")
    with pytest.raises(ResumeParseError, match="only be read on a Mac"):
        extract_text("resume.doc", b"\xd0\xcf\x11\xe0 legacy word file")


def test_bad_pdf():
    with pytest.raises(ResumeParseError):
        extract_text("resume.pdf", b"%PDF-1.4 garbage")
